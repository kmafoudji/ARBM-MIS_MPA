"""
Tests for the AS-IS import (core/asis_import) — design §12.

The workbooks are built in memory by `build_workbook`, which reproduces the
shape of the real files: thirteen sheets, annotated headers, and a note row
on some sheets only — which is what exercises the header-row detection.
"""
import io

import pytest
from openpyxl import Workbook
from rest_framework.test import APIClient

from tests.factories import (
    CountryFactory,
    CurrencyFactory,
    HubFactory,
    SdgFactory,
    SectorFactory,
    UserFactory,
)

ENDPOINT = "/api/import/asis/"

# Sheets whose row 1 carries a note, as in the real workbooks.
# 02a_envelope, 03_donors and 04_agencies start with their header.
SHEETS_WITH_NOTE = {
    "01_project",
    "02_financing_source",
    "05_project_partners",
    "06_toc_nodes",
    "07_indicators_logframe",
    "08_logframe_targets",
    "09_components",
    "10_activities",
    "11_milestones",
    "12_gadm_scope",
    "13_results_data",
}

HEADERS = {
    "01_project": [
        "official_reference_number", "name", "lead_country_iso3", "hub_code",
        "primary_sector", "sdgs", "we_category",
        "climate_marker", "risk_rating", "budget_amount", "currency",
        "reporting_frequency", "next_reporting_due", "lifecycle_stage",
        "start_date", "end_date",
    ],
    "02a_envelope": ["official_reference_number", "notes"],
    "02_financing_source": ["source (enum)", "instrument (enum)", "amount", "currency", "amount_usd", "note"],
    "03_donors": ["code", "name"],
    "04_agencies": ["code", "name", "agency_type_note"],
    "05_project_partners": ["agency_code", "role", "is_lead"],
    "06_toc_nodes": ["node_ref", "chain_level", "parent_ref", "statement", "assumptions"],
    "07_indicators_logframe": [
        "indicator_code", "chain_level", "sector", "name", "definition", "unit",
        "direction", "baseline_value", "baseline_year", "end_target_value",
        "end_target_date", "reporting_frequency", "toc_node_ref (informational)",
    ],
    "08_logframe_targets": ["indicator_code", "target_date", "target_value", "is_original_pad"],
    "09_components": ["code", "parent_code", "level", "sequence", "name"],
    "10_activities": [
        "activity_id", "sub_component_code", "name", "baseline_start", "baseline_end",
        "current_start", "current_end", "status", "budget_planned", "depends_on",
    ],
    "11_milestones": [
        "milestone_id", "activity_id", "name", "baseline_date", "current_date",
        "actual_date", "status",
    ],
    "12_gadm_scope": ["level", "admin1_name", "admin2_name", "admin3_name", "source_site"],
    "13_results_data": ["indicator_code", "period", "value", "source", "narrative"],
}

SHEET_ORDER = [
    "00_README", "98_Loader_Notes", "01_project", "02a_envelope", "02_financing_source",
    "03_donors", "04_agencies", "05_project_partners", "06_toc_nodes", "07_indicators_logframe",
    "08_logframe_targets", "09_components", "10_activities", "11_milestones",
    "12_gadm_scope", "13_results_data", "99_Parked",
]

# Emitted only when the test supplies rows for them: the base workbook keeps
# the 13-sheet shape.
OPTIONAL_SHEETS = {"06_toc_nodes"}

PROJECT_NAME = "Test project"


def base_rows(country_iso3, hub_code, sector_code):
    """A coherent, minimal project: the starting point of every test."""
    return {
        "01_project": [[
            "REF001", PROJECT_NAME, country_iso3, hub_code, sector_code,
            "1; 2; 5", "", "", "tbd", 1000, "USD", "quarterly",
            "2025-03-31", "LS013", "2025-01-01", "2026-12-31",
        ]],
        "02a_envelope": [["REF001", "Envelope note"]],
        "02_financing_source": [
            ["llf", "grant", 400, "USD", 400, "Grant share"],
            ["isdb_oc", "loan", 600, "USD", 600, "Loan share"],
        ],
        "03_donors": [["TESTDONOR", "Test donor"]],
        "04_agencies": [["TEST-AG", "Test agency", "government"]],
        "05_project_partners": [["TEST-AG", "lead", "Yes"]],
        "07_indicators_logframe": [[
            "REF001-IND-01", "output", sector_code, "Test indicator",
            "Definition", "Hectares", "increase", 0, 2024, 100, "2026-12-31",
            "annual", "REF001-OUT-1",
        ]],
        "08_logframe_targets": [],
        "09_components": [
            ["REF001-C1", "", "Component", 1, "Component 1"],
            ["REF001-C1.1", "REF001-C1", "Sub-component", 1, "Sub-component 1.1"],
        ],
        "10_activities": [[
            "REF001-A1", "REF001-C1.1", "Test activity",
            "2025-01-01", "2026-12-31", "2025-01-01", "2026-12-31",
            "not_started", 500, "",
        ]],
        "11_milestones": [],
        "12_gadm_scope": [],
        "13_results_data": [],
    }


def build_workbook(rows):
    """Serialise `rows` into an xlsx workbook and return an in-memory file."""
    book = Workbook()
    book.remove(book.active)
    for name in SHEET_ORDER:
        if name in OPTIONAL_SHEETS and name not in rows:
            continue
        sheet = book.create_sheet(name)
        if name not in HEADERS:
            sheet.append([f"Sheet {name} — ignored by the import."])
            continue
        if name in SHEETS_WITH_NOTE:
            sheet.append([f"Note for sheet {name}."])
        sheet.append(HEADERS[name])
        for row in rows.get(name, []):
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


@pytest.fixture
def reference_data(db):
    """Minimal reference tables: the parser resolves its FKs against the real database."""
    country = CountryFactory(iso3="TST", iso2="TS")
    hub = HubFactory(code="hub-test")
    sector = SectorFactory(code="sector-test")
    CurrencyFactory(code="USD")
    for number in (1, 2, 5):
        SdgFactory(number=number)
    return {"country": country, "hub": hub, "sector": sector}


@pytest.fixture
def rows(reference_data):
    return base_rows(
        reference_data["country"].iso3,
        reference_data["hub"].code,
        reference_data["sector"].code,
    )


@pytest.fixture
def auth_client(db):
    client = APIClient()
    client.force_authenticate(user=UserFactory())
    return client


def post(client, workbook, mode="validate", expected_sha256=None):
    payload = {"file": workbook, "mode": mode}
    if expected_sha256 is not None:
        payload["expected_sha256"] = expected_sha256
    return client.post(ENDPOINT, payload, format="multipart")


def messages(items):
    return " | ".join(item["message"] for item in items)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_valid_workbook_validates_without_errors(auth_client, rows):
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert response.data["errors"] == []
    assert response.data["project_ref"] == "REF001"
    assert response.data["project_exists"] is False
    assert response.data["summary"]["01_project"]["create"] == 1
    assert response.data["file_sha256"]


@pytest.mark.django_db
def test_enum_outside_model_choices_is_an_error(auth_client, rows):
    """Design §8: a value that is not a model literal blocks the import."""
    rows["10_activities"][0][7] = "Delayed"  # status
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 422
    assert any(
        error["sheet"] == "10_activities" and error["column"] == "status"
        for error in response.data["errors"]
    ), messages(response.data["errors"])


@pytest.mark.django_db
def test_unresolved_internal_reference_is_an_error(auth_client, rows):
    """An activity naming a sub-component absent from 09_components."""
    rows["10_activities"][0][1] = "REF001-C9.9"
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 422
    assert any(
        error["column"] == "sub_component_code" for error in response.data["errors"]
    ), messages(response.data["errors"])


@pytest.mark.django_db
def test_financing_must_reconcile_to_budget_amount(auth_client, rows):
    """VAL015: the financing rows must add back up to budget_amount."""
    rows["02_financing_source"][0][4] = 300  # amount_usd: 300 + 600 != 1000
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 422
    assert "VAL015" in messages(response.data["errors"])


@pytest.mark.django_db
def test_end_date_before_start_date_is_an_error(auth_client, rows):
    """VAL012."""
    rows["01_project"][0][15] = "2024-01-01"  # end_date < start_date
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 422
    assert "VAL012" in messages(response.data["errors"])


@pytest.mark.django_db
def test_unresolved_gadm_area_warns_and_does_not_block(auth_client, rows):
    """
    D-9: the workbooks describe levels the loaded GADM reference does not
    know. The gap is reported and the rest of the file is loaded.
    """
    rows["12_gadm_scope"] = [["ADM3", "Province", "District", "Chiefdom", "site S01"]]
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert any(
        warning["sheet"] == "12_gadm_scope" for warning in response.data["warnings"]
    ), messages(response.data["warnings"])
    assert "12_gadm_scope" not in response.data["summary"]


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_commit_writes_then_revalidation_reports_unchanged(auth_client, rows):
    """
    The proof that the upsert works (design §12, step 4): replaying the same
    file must change nothing.
    """
    from apps.project.models import Project
    from apps.workplan.models import Activity

    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200

    workbook.seek(0)
    commit = post(
        auth_client, workbook, mode="commit",
        expected_sha256=validation.data["file_sha256"],
    )
    assert commit.status_code == 200, commit.data
    assert commit.data["committed"] is True

    project = Project.objects.get(official_reference_number="REF001")
    assert project.pk == commit.data["project_id"]
    assert project.financial_envelope.total_amount_usd == 1000
    assert sorted(project.sdgs.values_list("number", flat=True)) == [1, 2, 5]
    assert project.logframe_rows.count() == 1
    assert Activity.objects.filter(sub_component__component__project=project).count() == 1

    workbook.seek(0)
    again = post(auth_client, workbook, mode="validate")
    assert again.status_code == 200
    assert again.data["project_exists"] is True
    # Every sheet, with no exception. Financing and milestones used to be
    # excluded here because replace-all could not report itself as unchanged;
    # syncing them by content is what removed that exception.
    actions = {change["action"] for change in again.data["changes"]}
    assert actions == {"unchanged"}, again.data["changes"]


@pytest.mark.django_db
def test_legacy_sdg_columns_are_folded_into_the_set(auth_client, rows):
    """ADR 0006: v0.x workbooks with primary_sdg + contributing_sdgs still load."""
    from apps.project.models import Project

    legacy_rows = dict(rows)
    legacy_rows["01_project"] = [list(rows["01_project"][0])]
    legacy_rows["01_project"][0][5] = ""  # the new `sdgs` column left empty
    legacy_headers = {name: list(cols) for name, cols in HEADERS.items()}
    legacy_headers["01_project"] = legacy_headers["01_project"] + ["primary_sdg", "contributing_sdgs"]
    legacy_rows["01_project"][0] += [1, "2; 5"]

    original = dict(HEADERS)
    HEADERS.update(legacy_headers)
    try:
        workbook = build_workbook(legacy_rows)
    finally:
        HEADERS.clear()
        HEADERS.update(original)

    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data
    assert any("legacy" in w["message"] for w in validation.data["warnings"]), \
        messages(validation.data["warnings"])
    workbook.seek(0)
    commit = post(auth_client, workbook, mode="commit",
                  expected_sha256=validation.data["file_sha256"])
    assert commit.status_code == 200, commit.data
    project = Project.objects.get(official_reference_number="REF001")
    assert sorted(project.sdgs.values_list("number", flat=True)) == [1, 2, 5]


@pytest.mark.django_db
def test_update_reports_a_field_level_diff(auth_client, rows):
    """Design §6: an `update` action carries the difference field by field."""
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])

    renamed = f"{PROJECT_NAME} renamed"
    rows["01_project"][0][1] = renamed
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200
    change = next(c for c in response.data["changes"] if c["sheet"] == "01_project")
    assert change["action"] == "update"
    diff = next(d for d in change["diffs"] if d["field"] == "name")
    assert diff["from"] == PROJECT_NAME
    assert diff["to"] == renamed


@pytest.mark.django_db
def test_blank_cell_never_clears_a_stored_value(auth_client, rows):
    """
    The parser's rule: an empty cell means "not supplied".

    The AS-IS workbooks leave classification columns empty deliberately;
    without this rule the first commit would erase what the interface
    already holds.
    """
    from apps.project.models import Project

    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])
    assert Project.objects.get(official_reference_number="REF001").risk_rating == "tbd"

    rows["01_project"][0][8] = ""  # risk_rating left empty
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200
    change = next(c for c in validation.data["changes"] if c["sheet"] == "01_project")
    assert change["action"] == "unchanged"

    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])
    assert Project.objects.get(official_reference_number="REF001").risk_rating == "tbd"


@pytest.mark.django_db
def test_activityless_milestone_gets_a_declared_placeholder(auth_client, rows):
    """
    D-10: Milestone.activity is a required FK. The placeholder activity is
    created, and its creation appears in the report.
    """
    from apps.workplan.models import Milestone

    rows["11_milestones"] = [[
        "REF001-M01", "", "Financing agreement signature",
        "2025-02-01", "", "", "pending",
    ]]
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data

    placeholder = [
        change for change in validation.data["changes"]
        if change["sheet"] == "10_activities" and "MILESTONES" in change["target"]
    ]
    assert placeholder and placeholder[0]["action"] == "create"

    workbook.seek(0)
    commit = post(auth_client, workbook, mode="commit",
                  expected_sha256=validation.data["file_sha256"])
    assert commit.status_code == 200, commit.data
    milestone = Milestone.objects.get(name="Financing agreement signature")
    assert milestone.activity.code.endswith("-A-MILESTONES")
    assert milestone.category == "contractual"


@pytest.mark.django_db
def test_a_stage_at_or_past_effective_activates_the_workspace(auth_client, rows):
    """
    The importer writes `lifecycle_stage` directly, so the SF-4 state machine
    never runs and never generates the workspace. `ProjectDetail` locks the
    Workplan and Results tabs on that row, which left everything the import
    loaded present in the database and unreachable in the interface.
    """
    from apps.project.models import Project, ProjectWorkspace

    workbook = build_workbook(rows)  # 01_project declares `LS013` (Implementing)
    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data
    planned = [c for c in validation.data["changes"] if c["sheet"] == "workspace"]
    assert planned and planned[0]["action"] == "create"

    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])

    project = Project.objects.get(official_reference_number="REF001")
    assert ProjectWorkspace.objects.filter(project=project).exists()

    # Idempotent: a second pass reports it as already activated.
    again = post(auth_client, build_workbook(rows))
    planned = [c for c in again.data["changes"] if c["sheet"] == "workspace"]
    assert planned and planned[0]["action"] == "unchanged"


@pytest.mark.django_db
def test_a_stage_before_effective_activates_nothing(auth_client, rows):
    """Below Effective the workspace is not implied and must not be invented."""
    rows["01_project"][0][13] = "LS009"  # lifecycle_stage: Appraisal
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert "workspace" not in response.data["summary"]


@pytest.mark.django_db
def test_setting_the_stage_directly_warns_that_no_audit_row_is_written(auth_client, rows):
    """
    A project loaded from a file has no transition history. Reconstructing it
    would invent dates, actors and justifications in the one table whose value
    is being immutable and true (POL-1.09), so the report says so instead.
    """
    from apps.project.models import ProjectStageTransition, Project

    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    assert any(
        w["column"] == "lifecycle_stage" for w in validation.data["warnings"]
    ), messages(validation.data["warnings"])

    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])
    project = Project.objects.get(official_reference_number="REF001")
    assert ProjectStageTransition.objects.filter(project=project).count() == 0


def test_the_written_field_sets_name_real_model_fields():
    """
    The replace-all warning subtracts these sets from the model to say what a
    confirm costs. A typo would silently inflate the list of losses, so the
    names are checked against the models. Nothing can prove the applier writes
    exactly these — that part stays a comment on the constants.
    """
    from apps.project.models import FinancingSource
    from apps.workplan.models import Milestone
    from core.asis_import.parser import (
        FINANCING_WRITTEN_FIELDS,
        MILESTONE_WRITTEN_FIELDS,
        _fields_reset_by_replacement,
    )

    for model, written in (
        (FinancingSource, FINANCING_WRITTEN_FIELDS),
        (Milestone, MILESTONE_WRITTEN_FIELDS),
    ):
        names = {f.name for f in model._meta.fields}
        assert written <= names, f"{model.__name__}: {written - names}"

    # The two the design missed until 21 August 2026.
    assert "donor" in _fields_reset_by_replacement(FinancingSource, FINANCING_WRITTEN_FIELDS)
    assert "is_gate" in _fields_reset_by_replacement(Milestone, MILESTONE_WRITTEN_FIELDS)


MILESTONE_ROW = [
    "REF001-M01", "", "Financing agreement signature",
    "2025-02-01", "", "", "pending",
]


def commit_once(client, rows):
    """Validate then commit, returning the committed report."""
    workbook = build_workbook(rows)
    validation = post(client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data
    workbook.seek(0)
    committed = post(client, workbook, mode="commit",
                     expected_sha256=validation.data["file_sha256"])
    assert committed.status_code == 200, committed.data
    return committed


@pytest.mark.django_db
def test_an_unchanged_reload_warns_about_nothing_and_touches_nothing(auth_client, rows):
    """
    The point of syncing by content: a file that has not changed costs
    nothing. No warning, no write, and the stored rows keep their identity.
    """
    from apps.project.models import FinancingSource, Project
    from apps.workplan.models import Milestone

    rows["11_milestones"] = [list(MILESTONE_ROW)]
    commit_once(auth_client, rows)

    project = Project.objects.get(official_reference_number="REF001")
    before_financing = set(
        FinancingSource.objects.filter(envelope__project=project).values_list("pk", flat=True)
    )
    before_milestones = set(
        Milestone.objects.filter(
            activity__sub_component__component__project=project
        ).values_list("pk", flat=True)
    )

    again = post(auth_client, build_workbook(rows))
    assert again.status_code == 200
    text = messages(again.data["warnings"])
    assert "will be deleted" not in text

    commit_once(auth_client, rows)
    assert before_financing == set(
        FinancingSource.objects.filter(envelope__project=project).values_list("pk", flat=True)
    )
    assert before_milestones == set(
        Milestone.objects.filter(
            activity__sub_component__component__project=project
        ).values_list("pk", flat=True)
    )


@pytest.mark.django_db
def test_hand_edited_fields_survive_an_unchanged_reload(auth_client, rows):
    """
    `donor` and `is_gate` have no column in the workbook. Under replace-all
    they were cleared on every load; rows that do not change are now left
    alone, so they survive. `is_gate` is a control, not just a value — it is
    what stops an activity reaching 100% before its milestone (RG-5.2).
    """
    from apps.project.models import FinancingSource, Project
    from apps.reference.models import Donor
    from apps.workplan.models import Milestone

    rows["11_milestones"] = [list(MILESTONE_ROW)]
    commit_once(auth_client, rows)
    project = Project.objects.get(official_reference_number="REF001")

    donor = Donor.objects.create(code="handpicked", name="Entered through the interface")
    financing = FinancingSource.objects.filter(envelope__project=project).first()
    financing.donor = donor
    financing.save()
    milestone = Milestone.objects.filter(
        activity__sub_component__component__project=project
    ).first()
    milestone.is_gate = True
    milestone.evidence_url = "https://example.test/evidence.pdf"
    milestone.save()

    commit_once(auth_client, rows)

    financing.refresh_from_db()
    milestone.refresh_from_db()
    assert financing.donor == donor
    assert milestone.is_gate is True
    assert milestone.evidence_url == "https://example.test/evidence.pdf"


@pytest.mark.django_db
def test_a_changed_row_is_one_delete_and_one_create(auth_client, rows):
    """
    Without a key there is no way to recognise a changed row as the same row,
    so it is removed and written again — but only that row. The others stay
    untouched, which is the whole difference from replace-all.
    """
    rows["02_financing_source"] = [
        ["llf", "grant", 400, "USD", 400, "Grant share"],
        ["isdb_oc", "loan", 600, "USD", 600, "Loan share"],
    ]
    commit_once(auth_client, rows)

    # The note changes; the amounts still reconcile, so VAL015 stays happy.
    rows["02_financing_source"][1][5] = "Loan share — renegotiated"
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    counts = response.data["summary"]["02_financing_source"]
    assert counts["create"] == 1
    assert counts["delete"] == 1
    assert counts["unchanged"] == 1
    text = messages(response.data["warnings"])
    assert "is_gate" not in text  # the milestone warning must not fire here
    assert "1 financing row(s) will be deleted and 1 created" in text


@pytest.mark.django_db
def test_replacing_milestones_warns_when_alerts_hang_off_them(auth_client, rows):
    """
    WorkplanAlert.milestone is CASCADE. The design asserted nothing referenced
    Milestone; it was wrong, so the alerts are counted and reported.
    """
    from apps.workplan.models import Milestone, WorkplanAlert

    rows["11_milestones"] = [[
        "REF001-M01", "", "Financing agreement signature",
        "2025-02-01", "", "", "pending",
    ]]
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    workbook.seek(0)
    post(auth_client, workbook, mode="commit",
         expected_sha256=validation.data["file_sha256"])

    milestone = Milestone.objects.get(name="Financing agreement signature")
    WorkplanAlert.objects.create(
        project=milestone.activity.project,
        activity=milestone.activity,
        milestone=milestone,
        alert_type="milestone_t30",
        message="Milestone due in 30 days.",
    )

    # Unchanged file: the milestone is not deleted, so nothing is at risk.
    quiet = post(auth_client, build_workbook(rows))
    assert "workplan alert(s)" not in messages(quiet.data["warnings"])

    # Change the milestone: now it is deleted and recreated, and the alert
    # goes with it (WorkplanAlert.milestone is CASCADE).
    rows["11_milestones"][0][6] = "achieved"
    again = post(auth_client, build_workbook(rows))
    assert again.status_code == 200
    assert "workplan alert(s)" in messages(again.data["warnings"])


@pytest.mark.django_db
def test_an_acronym_named_reference_row_is_flagged_as_a_near_duplicate(auth_client, rows):
    """
    The case a name-only comparison missed on the real data: the reference
    table held `knarda` named "KNARDA" while NGA1007 declared `NGA-KNARDA`
    named "Kano State Agricultural and Rural Development Authority". Same
    body, no shared name — the signal is in the code.
    """
    from apps.reference.models import ImplementingAgency

    ImplementingAgency.objects.create(
        code="knarda", name="KNARDA", agency_type="national_agency"
    )
    rows["04_agencies"] = [[
        "NGA-KNARDA", "Kano State Agricultural and Rural Development Authority", "government",
    ]]
    rows["05_project_partners"] = [["NGA-KNARDA", "lead", "Yes"]]

    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    flagged = [
        w for w in response.data["warnings"]
        if w["sheet"] == "04_agencies" and w["column"] == "code"
    ]
    assert flagged, messages(response.data["warnings"])
    assert "knarda" in flagged[0]["message"]
    # A warning, never a merge: the reference tables are governed outside
    # this import and rows pointing at them would have to be repointed.
    assert response.data["errors"] == []


@pytest.mark.django_db
def test_a_short_shared_code_segment_is_not_a_near_duplicate(auth_client, rows):
    """
    `NGA-PMU` and `SLE-PMU` share `PMU` and are different bodies. Below the
    length floor nothing fires — a warning that cries wolf is a warning
    nobody reads.
    """
    from apps.reference.models import ImplementingAgency

    ImplementingAgency.objects.create(
        code="SLE-PMU", name="SL-RVCP PMU", agency_type="national_agency"
    )
    rows["04_agencies"] = [["NGA-PMU", "KSADP PMU", "government"]]
    rows["05_project_partners"] = [["NGA-PMU", "lead", "Yes"]]

    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert not [
        w for w in response.data["warnings"]
        if w["sheet"] == "04_agencies" and w["column"] == "code"
    ], messages(response.data["warnings"])


@pytest.mark.django_db
def test_existing_donor_is_matched_case_insensitively(auth_client, rows):
    """
    The database holds `isdb`, the workbooks write `ISDB`. A literal match
    would duplicate a reference row that already exists.
    """
    from apps.reference.models import Donor

    Donor.objects.create(code="testdonor", name="Donor already in the reference table")
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200
    change = next(c for c in response.data["changes"] if c["sheet"] == "03_donors")
    assert change["action"] == "unchanged"
    assert Donor.objects.filter(code__iexact="testdonor").count() == 1


@pytest.mark.django_db
def test_a_new_indicator_code_is_catalogued_as_project_specific(auth_client, rows):
    """
    A code the catalogue does not have belongs to the project that brought it,
    not to the institutional catalogue: the loader says so on the row instead
    of leaving the model default to mean it by accident.
    """
    from apps.results.models import Indicator

    validation = post(auth_client, build_workbook(rows))
    change = next(
        c for c in validation.data["changes"]
        if c["sheet"] == "07_indicators_logframe" and c["action"] == "create"
    )
    assert "Project-specific" in change["detail"]

    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    workbook.seek(0)
    commit = post(auth_client, workbook, mode="commit",
                  expected_sha256=validation.data["file_sha256"])
    assert commit.status_code == 200, commit.data
    assert Indicator.objects.get(code="REF001-IND-01").indicator_type == "project_specific"


@pytest.mark.django_db
def test_an_indicator_already_in_the_catalogue_keeps_its_type(auth_client, rows, reference_data):
    """POL-2.01: a project file does not rewrite a catalogue the LLFMU governs."""
    from apps.results.models import Indicator

    Indicator.objects.create(
        code="REF001-IND-01",
        sector=reference_data["sector"],
        name="Institutional wording",
        indicator_type="percentage",
        direction="increase",
        definition="A definition.",
        unit="Percent",
    )
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    workbook.seek(0)
    commit = post(auth_client, workbook, mode="commit",
                  expected_sha256=validation.data["file_sha256"])
    assert commit.status_code == 200, commit.data

    indicator = Indicator.objects.get(code="REF001-IND-01")
    assert indicator.indicator_type == "percentage"
    assert indicator.name == "Institutional wording"


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_commit_with_a_mismatched_hash_returns_409(auth_client, rows):
    """
    D-4: nothing is stored between the two calls, so the hash echo is what
    stops someone validating file A and confirming file B.
    """
    from apps.project.models import Project

    response = post(
        auth_client, build_workbook(rows), mode="commit",
        expected_sha256="0" * 64,
    )
    assert response.status_code == 409
    assert not Project.objects.filter(official_reference_number="REF001").exists()


@pytest.mark.django_db
def test_commit_without_a_hash_is_refused(auth_client, rows):
    response = post(auth_client, build_workbook(rows), mode="commit")
    assert response.status_code == 400


@pytest.mark.django_db
def test_a_file_that_is_not_a_zip_is_refused(auth_client):
    """Design §11: the signature is checked on the content, not on the claim."""
    fake = io.BytesIO(b"This is not a workbook.")
    fake.name = "fake.xlsx"
    response = auth_client.post(
        ENDPOINT, {"file": fake, "mode": "validate"}, format="multipart"
    )
    assert response.status_code == 400


@pytest.mark.django_db
def test_a_workbook_missing_sheets_is_refused(auth_client):
    book = Workbook()
    book.remove(book.active)
    book.create_sheet("01_project")
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    response = post(auth_client, buffer)
    assert response.status_code == 422
    assert "Sheets missing" in messages(response.data["errors"])


@pytest.mark.django_db
def test_the_endpoint_requires_authentication(rows):
    response = APIClient().post(
        ENDPOINT, {"file": build_workbook(rows), "mode": "validate"}, format="multipart"
    )
    assert response.status_code in (401, 403)


@pytest.mark.django_db
def test_unknown_project_column_is_reported(auth_client, rows):
    """
    A header the import does not read (for instance a column retired by the
    classification rework) is a warning: its values are silently ignored
    otherwise. Known headers raise no such warning.
    """
    workbook = build_workbook(rows)
    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data
    assert not any("not read by the import" in w["message"] for w in validation.data["warnings"]), \
        messages(validation.data["warnings"])

    old_rows = {name: [list(r) for r in rs] for name, rs in rows.items()}
    old_headers = {name: list(cols) for name, cols in HEADERS.items()}
    old_headers["01_project"] = old_headers["01_project"] + ["gender_marker", "fragility_status"]
    old_rows["01_project"][0] += ["1", "fcv"]

    original = dict(HEADERS)
    HEADERS.update(old_headers)
    try:
        workbook = build_workbook(old_rows)
    finally:
        HEADERS.clear()
        HEADERS.update(original)

    validation = post(auth_client, workbook, mode="validate")
    assert validation.status_code == 200, validation.data
    warning = next(
        (w for w in validation.data["warnings"] if "not read by the import" in w["message"]), None
    )
    assert warning is not None, messages(validation.data["warnings"])
    assert warning["sheet"] == "01_project"
    assert "fragility_status" in warning["message"] and "gender_marker" in warning["message"]


# ---------------------------------------------------------------------------
# 06_toc_nodes (optional sheet, decision 0010)
# ---------------------------------------------------------------------------


def toc_rows():
    """
    A complete 5-level tree written top-down, as the PAD reads it, with the
    fan-in the model cannot hold as `parent`: two activities feed one
    output, and two outputs feed one immediate outcome.
    """
    return [
        ["REF001-IMP-1", "Ultimate Outcome", "", "Impact statement", "Stable economy"],
        ["REF001-OC-1", "Intermediate Outcome", "REF001-IMP-1", "Intermediate 1", ""],
        ["REF001-IO-1", "Immediate Outcome", "REF001-OC-1", "Immediate 1", "Farmers adopt"],
        ["REF001-OUT-1", "Output", "REF001-IO-1", "Output 1", ""],
        ["REF001-OUT-2", "Output", "REF001-IO-1", "Output 2", ""],
        ["REF001-ACT-1", "activity", "REF001-OUT-1", "Activity 1", ""],
        ["REF001-ACT-2", "activity", "REF001-OUT-1", "Activity 2", ""],
        ["REF001-ACT-3", "activity", "REF001-OUT-2", "Activity 3", ""],
    ]


@pytest.mark.django_db
def test_without_the_toc_sheet_nothing_about_the_toc_is_planned(auth_client, rows):
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert "06_toc_nodes" not in response.data["summary"]
    assert "06_toc_nodes" not in messages(response.data["warnings"])


@pytest.mark.django_db
def test_a_complete_toc_sheet_is_planned_bottom_up_with_cross_pathways(auth_client, rows):
    rows["06_toc_nodes"] = toc_rows()
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert response.data["errors"] == []
    counts = response.data["summary"]["06_toc_nodes"]
    assert counts["create"] == 8  # 7 nodes + the ultimate outcome field
    changes = [c for c in response.data["changes"] if c["sheet"] == "06_toc_nodes"]
    targets = [c["target"] for c in changes]
    # Activities before outputs before outcomes: parents exist when needed.
    assert targets.index("REF001-ACT-1") < targets.index("REF001-OUT-1")
    assert targets.index("REF001-OUT-2") < targets.index("REF001-IO-1")
    fan_in = next(c for c in changes if c["target"] == "REF001-OUT-1")
    assert "2 contributors, 1 as cross-pathway" in fan_in["detail"]


@pytest.mark.django_db
def test_committing_the_toc_writes_the_tree_the_model_way(auth_client, rows):
    from apps.project.models import Project
    from apps.results.models import ToCNode

    rows["06_toc_nodes"] = toc_rows()
    rows["07_indicators_logframe"][0][12] = "REF001-OUT-1"  # toc_node_ref
    commit_once(auth_client, rows)

    project = Project.objects.get(official_reference_number="REF001")
    toc = project.theory_of_change
    assert toc.ultimate_outcome == "Impact statement"
    nodes = {n.statement: n for n in ToCNode.objects.filter(toc=toc)}
    assert len(nodes) == 7
    # Inverted: the model's parent is the first contributor in file order.
    assert nodes["Activity 1"].parent is None
    assert nodes["Output 1"].parent == nodes["Activity 1"]
    assert nodes["Output 2"].parent == nodes["Activity 3"]
    assert nodes["Immediate 1"].parent == nodes["Output 1"]
    assert nodes["Intermediate 1"].parent == nodes["Immediate 1"]
    assert nodes["Immediate 1"].assumptions == "Farmers adopt"
    # The other contributors are cross-pathways, contributor -> target.
    assert list(nodes["Activity 2"].cross_pathways.all()) == [nodes["Output 1"]]
    assert list(nodes["Output 2"].cross_pathways.all()) == [nodes["Immediate 1"]]
    assert nodes["Activity 1"].cross_pathways.count() == 0
    # Codes come from create_toc_node, as on the ToC page.
    assert nodes["Activity 1"].code == "A"
    assert nodes["Output 1"].code == "A.1"
    assert nodes["Immediate 1"].code == "A.1.1"
    # 07.toc_node_ref becomes the logframe link.
    assert nodes["Output 1"].logframe_row == project.logframe_rows.get()
    assert nodes["Output 2"].logframe_row is None


@pytest.mark.django_db
def test_an_identical_toc_reload_is_unchanged_and_writes_nothing(auth_client, rows):
    from apps.results.models import ToCNode

    rows["06_toc_nodes"] = toc_rows()
    commit_once(auth_client, rows)
    before = sorted(ToCNode.objects.values_list("pk", flat=True))

    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    counts = response.data["summary"]["06_toc_nodes"]
    assert counts["unchanged"] == 8
    assert counts["create"] == 0 and counts["delete"] == 0
    assert sorted(ToCNode.objects.values_list("pk", flat=True)) == before


@pytest.mark.django_db
def test_a_changed_toc_replaces_the_whole_tree_visibly(auth_client, rows):
    from apps.results.models import ToCNode

    rows["06_toc_nodes"] = toc_rows()
    commit_once(auth_client, rows)

    rows["06_toc_nodes"][3][3] = "Output 1 — reworded"
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    counts = response.data["summary"]["06_toc_nodes"]
    assert counts["delete"] == 7
    assert counts["create"] == 7
    assert counts["unchanged"] == 1  # the ultimate outcome field
    assert "replaced as a whole" in messages(response.data["warnings"])

    commit_once(auth_client, rows)
    statements = set(ToCNode.objects.values_list("statement", flat=True))
    assert "Output 1 — reworded" in statements and "Output 1" not in statements
    assert ToCNode.objects.count() == 7


@pytest.mark.django_db
def test_an_incomplete_tree_warns_and_skips_the_sheet(auth_client, rows):
    """The shape of the legacy PADs: no activities, no immediate outcomes."""
    rows["06_toc_nodes"] = [
        ["REF001-IMP-1", "Ultimate Outcome", "", "Impact statement", ""],
        ["REF001-OC-1", "Intermediate Outcome", "REF001-IMP-1", "Intermediate 1", ""],
        ["REF001-OUT-1", "Output", "REF001-OC-1", "Output 1", ""],
    ]
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert response.data["errors"] == []
    assert "06_toc_nodes" not in response.data["summary"]
    text = messages(response.data["warnings"])
    # OUT-1 hangs straight off an intermediate outcome: a skipped tier, and
    # an output with no activity under it. Both are warnings, not errors.
    assert "REF001-OUT-1 (output): parent_ref REF001-OC-1 is intermediate_outcome, more than one level up" in text
    assert "1 node(s) have no contributor and 1 skip a level" in text

    rows["06_toc_nodes"] = [
        ["REF001-IMP-1", "Ultimate Outcome", "", "Impact statement", ""],
        ["REF001-OC-1", "Intermediate Outcome", "REF001-IMP-1", "Intermediate 1", ""],
        ["REF001-IO-1", "Immediate Outcome", "REF001-OC-1", "Immediate 1", ""],
        ["REF001-OUT-1", "Output", "REF001-IO-1", "Output 1", ""],
    ]
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 200, response.data
    assert response.data["errors"] == []
    assert "06_toc_nodes" not in response.data["summary"]
    text = messages(response.data["warnings"])
    assert "REF001-OUT-1 (output) has no contributing activity" in text
    assert "Sheet skipped: 1 node(s) have no contributor and 0 skip a level" in text
    # The rest of the workbook is planned as usual.
    assert response.data["summary"]["01_project"]["create"] == 1


@pytest.mark.django_db
def test_toc_format_defects_are_errors(auth_client, rows):
    rows["06_toc_nodes"] = toc_rows() + [
        ["REF001-ACT-9", "activity", "REF001-NOPE", "Dangling", ""],
        ["REF001-ACT-1", "activity", "REF001-OUT-1", "Duplicate ref", ""],
        ["REF001-X", "impact", "REF001-IMP-1", "Unknown level", ""],
        ["REF001-OUT-3", "Output", "REF001-ACT-3", "Upside down", ""],
    ]
    response = post(auth_client, build_workbook(rows))
    assert response.status_code == 422, response.data
    text = messages(response.data["errors"])
    assert "parent_ref 'REF001-NOPE' is not in the sheet" in text
    assert "REF001-OUT-3 (output): parent_ref REF001-ACT-3 is activity, which is not above it" in text
    assert "REF001-ACT-1 appears more than once" in text
    assert "'impact' is not a value the model accepts" in text
    assert "06_toc_nodes" not in response.data["summary"]
